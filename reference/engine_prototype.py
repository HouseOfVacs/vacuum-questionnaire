"""Reference prototype of the HOV Vacuum Recommendation Questionnaire logic (approved Oct 14, 2026).
Reads HOV_Questionnaire_Catalog_v1.xlsx and runs the seven approved test households.
Claude Code: port this to JavaScript as-is; the printed results are the expected test outputs.
Add-ons and constants below mirror the Add-ons and Settings tabs of the catalog sheet."""
import openpyxl

SOFT = {"name": "Soft bristle brush roller (5010GE)", "url": "https://www.thehouseofvacs.com/store/p/sebo-brush-roller-12-5010ge-soft-bristle"}
SOFT_X8 = {"name": "Soft bristle brush roller (5290GE)", "url": None}  # not on store yet -> "ask us"
KOMBI = {"name": "Kombi Floor Nozzle Premium (7260GS)", "url": "https://www.thehouseofvacs.com/store/p/sebo-kombi-floor-nozzle-premium"}

# add-ons that turn a flag to OK: {model: {flag: addon}}
ADDONS = {
    "sebo-x4": {"shag": SOFT}, "sebo-x4-boost": {"shag": SOFT}, "sebo-x7": {"shag": SOFT},
    "sebo-x8": {"shag": SOFT_X8}, "sebo-felix": {"shag": SOFT}, "sebo-dart": {"shag": SOFT},
    "sebo-d4": {"wool": KOMBI}, "sebo-d1": {"wool": KOMBI}, "sebo-e3": {"wool": KOMBI},
    "sebo-e2": {"wool": KOMBI}, "sebo-k3": {"wool": KOMBI},
    "sebo-k2-turbo": {"wool": KOMBI, "shag": KOMBI},
}
PRICE_BANDS = {"good_max": 800, "better_max": 1000}   # Good < 800 <= Better < 1000 <= Best
MIN_FIT = 3.0          # below this a model is not shown
TIE_WINDOW = 0.1       # scores within this are a tie -> house boost decides
CAUTION_PENALTY = 0.3
HEPA_BONUS = 0.2

def load(path):
    """Read the Models tab of HOV_Questionnaire_Catalog_v1.xlsx (the Google Sheet export)."""
    ws = openpyxl.load_workbook(path)["Models"]
    head = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(head)}
    models = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        g = lambda h: row[col[h]]
        if not g("ID") or g("Active") != "Yes":
            continue
        style = {"Cordless stick": "Cordless (stick)", "Cordless upright": "Cordless (upright)"}.get(g("Style"), g("Style"))
        name = g("Model name")
        name = name[len(g("Brand")) + 1:] if name.startswith(g("Brand") + " ") else name
        models.append(dict(id=g("ID"), brand=g("Brand"), name=name, price=float(g("Price")), url=g("Store link"),
                           style=style, role=g("Cordless role") or "—",
                           carpet=g("Carpet"), bare=g("Bare floor"), pet=g("Pet hair"),
                           stairs=g("Stairs & above-floor"), ease=g("Ease of handling"),
                           flags={"wool": g("Wool / hand-sewn"), "shag": g("Shag / high pile"), "looped": g("Looped / berber")},
                           pick=(g("Owner's Pick") == "Yes"), boost=g("House favorite boost") or 0,
                           hepa=(g("HEPA") == "Yes")))
    return models

WOOL = {"Wool (any kind)", "Persian (Hand Sewn)", "Jute"}
SHAG = {"Synthetic Cut High Pile (Shag & Frieze)", "Looped High Pile"}
LOOPED = {"Textured (looped & cut)", "Looped Low Pile", "Looped Medium Pile", "Looped High Pile"}
DELICATE = {"Hardwood", "Bamboo", "Laminate Hardwood"}

def recommend(models, a):
    rugs = set() if a["mix"] == "none" else set(a["carpets"])
    need = {"wool": bool(rugs & WOOL), "shag": bool(rugs & SHAG), "looped": bool(rugs & LOOPED)}
    styles = a["styles"]

    def style_ok(m):
        if "any" in styles:
            return m["role"] != "Quick pickups"          # sticks only when asked for
        if m["style"] == "Upright" and "upright" in styles: return True
        if m["style"] == "Canister" and "canister" in styles: return True
        if m["style"].startswith("Cordless") and "cordless" in styles:
            want = "Main vacuum" if a.get("cordless_role") == "main" else "Quick pickups"
            return m["role"] == want
        return False

    # weights
    cw = {"carpet": 3, "half": 2, "bare": 1, "none": 0}[a["mix"]]
    bw = {"carpet": 1, "half": 2, "bare": 3, "none": 3}[a["mix"]] + (0.5 if set(a["floors"]) & DELICATE else 0)
    pw = {"none": 0, "low": 0.5, "moderate": 1.5, "heavy": 3}[a["pets"]]
    sw = {"no": 0, "carpeted": 1.5, "bare": 1}[a["stairs"]]
    s = a["slider"]                                   # 1 = lighter ... 5 = deepest clean
    ew = (5 - s) * 0.75 + {"daily": 1, "several": 0.5, "weekly": 0, "less": 0}[a["freq"]]
    mult = 0.8 + 0.1 * s
    cw, bw, pw = cw * mult, bw * mult, pw * mult

    def evaluate(m):
        addons, notes, pen = [], [], 0.0
        for f, needed in need.items():
            if not needed: continue
            flag = m["flags"][f]
            fix = ADDONS.get(m["id"], {}).get(f)
            if flag != "OK" and fix:
                if fix not in addons: addons.append(fix)
                continue
            if flag == "Avoid": return None
            if flag == "Caution":
                pen += CAUTION_PENALTY; notes.append(f)
        if a["limit"] and m["ease"] < 4: return None
        tw = cw + bw + pw + sw + ew
        score = (cw*m["carpet"] + bw*m["bare"] + pw*m["pet"] + sw*m["stairs"] + ew*m["ease"]) / tw
        score += (HEPA_BONUS if a["allergies"] and m["hepa"] else 0) - pen
        return dict(m, score=round(score, 2), addons=addons, cautions=notes)

    pool = [r for m in models if style_ok(m) and (r := evaluate(m)) and r["score"] >= MIN_FIT]
    def best_of(cands):
        """Highest fit; scores within TIE_WINDOW of the leader tie -> higher boost, then lower price."""
        if not cands: return None
        lead = max(c["score"] for c in cands)
        tied = [c for c in cands if lead - c["score"] <= TIE_WINDOW + 1e-9]
        return min(tied, key=lambda c: (-c["boost"], c["price"]))

    band = lambda p: "Good" if p < PRICE_BANDS["good_max"] else "Better" if p < PRICE_BANDS["better_max"] else "Best"
    chosen = [w for b in ("Good", "Better", "Best") if (w := best_of([r for r in pool if band(r["price"]) == b]))]
    pool.sort(key=lambda r: (-r["score"], -r["boost"], r["price"]))
    for r in pool:                                     # fill empty bands with next-best fits
        if len(chosen) >= 3: break
        if r not in chosen: chosen.append(r)
    chosen.sort(key=lambda r: r["price"])
    if len(chosen) >= 3:
        out = [dict(c, slot=l) for c, l in zip(chosen, ["Good", "Better", "Best"])]
        top = best_of(out)
        top["best_match"] = True                        # "Best match for your home"
    else:
        first = best_of(chosen)
        ranked = [first] + [c for c in chosen if c is not first]
        out = [dict(c, slot=l) for c, l in zip(ranked, ["Top pick", "Also great"])]

    if styles == {"cordless"} and a.get("cordless_role") == "main" and len(out) < 3:        # easy-handling corded alternative
        alt = recommend(models, dict(a, styles={"upright", "canister"}))
        cands = [x for x in alt["pool"] if x["ease"] >= 4]
        if cands:
            out.append(dict(best_of(cands), slot="Corded alternative", best_match=False))
    fallback = len([o for o in out if o["slot"] != "Corded alternative"]) < 2
    return {"picks": out, "fallback": fallback, "pool": pool}

HOUSEHOLDS = {
 "A. Carpet family, heavy-shedding dog": dict(styles={"any"}, mix="carpet", floors=["Vinyl Plank"], carpets=["Synthetic Cut Medium Pile (Frieze)"], pets="heavy", freq="several", stairs="carpeted", allergies=True, slider=3, limit=False),
 "B. Hardwood home with wool rugs, no pets": dict(styles={"upright", "canister"}, mix="bare", floors=["Hardwood"], carpets=["Wool (any kind)"], pets="none", freq="weekly", stairs="no", allergies=False, slider=3, limit=False),
 "C. Arthritis, half and half, cat": dict(styles={"any"}, mix="half", floors=["Laminate Hardwood"], carpets=["Synthetic Cut Low Pile"], pets="moderate", freq="weekly", stairs="no", allergies=False, slider=2, limit=True),
 "D. Cordless main vacuum, mostly carpet": dict(styles={"cordless"}, cordless_role="main", mix="carpet", floors=["Tile"], carpets=["Synthetic Cut Low Pile"], pets="moderate", freq="several", stairs="carpeted", allergies=False, slider=3, limit=False),
 "E. Shag + heavy shedder, uprights only, deepest clean": dict(styles={"upright"}, mix="carpet", floors=[], carpets=["Synthetic Cut High Pile (Shag & Frieze)"], pets="heavy", freq="weekly", stairs="no", allergies=False, slider=5, limit=False),
 "F. No carpet (tile + LVP), cat, quick-pickup cordless": dict(styles={"cordless"}, cordless_role="quick", mix="none", floors=["Tile", "Vinyl Plank"], carpets=[], pets="moderate", freq="daily", stairs="no", allergies=True, slider=2, limit=False),
 "G. Berber + stairs, allergies, no preference, lighter": dict(styles={"any"}, mix="half", floors=["Hardwood"], carpets=["Looped Low Pile"], pets="low", freq="weekly", stairs="carpeted", allergies=True, slider=2, limit=False),
}

if __name__ == "__main__":
    models = load("HOV_Questionnaire_Catalog_v1.xlsx")
    for name, a in HOUSEHOLDS.items():
        res = recommend(models, a)
        print("\n" + name + ("   [FALLBACK: worth a conversation]" if res["fallback"] else ""))
        for p in res["picks"]:
            extra = ""
            if p["addons"]: extra += "  + " + ", ".join(x["name"] for x in p["addons"])
            if p["cautions"]: extra += "  (caution: " + ", ".join(p["cautions"]) + ")"
            if p.get("best_match"): extra += "  << BEST MATCH FOR YOUR HOME"
            print(f"  {p['slot']:<19}{p['name']:<36}${p['price']:>8,.2f}  fit {p['score']:.2f}{extra}")
